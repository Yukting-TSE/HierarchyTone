#!/usr/bin/env python3
"""Generate accessible PPT for visually impaired teenagers experiment briefing."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# High-contrast palette for partial sight & projection
BG_DARK = RGBColor(0x1A, 0x1A, 0x2E)
ACCENT = RGBColor(0xFF, 0xD6, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xE8, 0xE8, 0xF0)
SUBTLE = RGBColor(0xB8, 0xB8, 0xCC)

SLIDES = [
    {
        "title": "欢迎参加听觉实验",
        "bullets": [
            "今天我们会一起体验两个有趣的声音实验",
            "全程用耳朵听、用手操作，不需要看屏幕",
            "有任何问题或不舒服，随时告诉我们",
        ],
        "notes": (
            "大家好！欢迎来参加今天的听觉实验活动。"
            "今天我们会一起做两个和声音有关的小实验，全程主要靠耳朵听、用键盘或鼠标操作，"
            "不需要看电脑屏幕。如果中途有任何问题、听不懂、或者觉得不舒服，"
            "请随时举手或告诉我们，我们会马上帮你。"
        ),
    },
    {
        "title": "实验一共两个部分",
        "bullets": [
            "第一部分：听着声音走迷宫",
            "第二部分：体验新的读屏软件",
            "结束后简单聊聊感受，还有小零食饮料",
        ],
        "notes": (
            "今天的活动分成两个部分。"
            "第一部分是「听着声音走迷宫」——你会在电脑上用方向键控制角色，靠声音找到出口。"
            "第二部分是「体验新的读屏软件」——我们会测试一种用不同音高告诉你信息重要程度的新方式。"
            "两个部分都做完之后，我们会口头聊聊你的使用感受，"
            "还会准备一些小零食和饮料，感谢你的参与。"
        ),
    },
    {
        "title": "第一部分：听着声音走迷宫",
        "bullets": [
            "在电脑迷宫里，用耳朵判断方向",
            "有「简单」和「复杂」两种难度",
            "每种难度各有 3 个迷宫，共 6 个",
        ],
        "notes": (
            "现在介绍第一部分：听着声音走迷宫。"
            "你会戴上专门的头戴式耳机，坐在电脑前。"
            "迷宫里你看不到画面，要靠耳朵听声音来判断该往哪个方向走。"
            "迷宫分「简单」和「复杂」两种难度。"
            "每种难度下面又有 3 种不同类型的迷宫，所以一共要通关 6 条迷宫。"
        ),
    },
    {
        "title": "迷宫 ①：听着音高随意走",
        "bullets": [
            "第一个迷宫：盲走模式",
            "根据音高的变化判断方向",
            "熟悉用声音在迷宫里移动的感觉",
        ],
        "notes": (
            "第一个迷宫是最基础的「盲走」模式。"
            "你会听到不同音高的声音，根据声音的变化来判断该往哪个方向走。"
            "这个迷宫主要是让你熟悉：不用眼睛，只靠耳朵在迷宫里移动是什么感觉。"
        ),
    },
    {
        "title": "迷宫 ②：沿着旋律的路线走",
        "bullets": [
            "迷宫的路线变成了一段旋律",
            "先记住旋律的音高变化",
            "再沿着这段旋律的声音走到终点",
        ],
        "notes": (
            "第二个迷宫更有挑战性：迷宫的正确路线被编成了一段旋律。"
            "你需要先听并记住这段旋律的音高变化——哪段音高、哪段音低，"
            "然后在走迷宫的时候，沿着这段旋律的声音方向前进，直到到达终点。"
            "就像跟着一首熟悉的歌的旋律走一样。"
        ),
    },
    {
        "title": "迷宫 ③：走错会有提醒",
        "bullets": [
            "和第二个迷宫一样，路线是旋律",
            "走错了会听到「错误噪音」提醒",
            "帮你及时知道自己偏离了正确路线",
        ],
        "notes": (
            "第三个迷宫在第二个的基础上加了一个新功能：错误提醒。"
            "路线依然是旋律，但如果你走错了方向，"
            "会立刻听到一种不好听的「错误噪音」，提醒你偏离了正确路线。"
            "这样你可以马上调整方向，重新跟着旋律走。"
        ),
    },
    {
        "title": "6 条迷宫怎么安排？",
        "bullets": [
            "简单难度：迷宫 ① ② ③ 各 1 个",
            "复杂难度：迷宫 ① ② ③ 各 1 个",
            "2 种难度 × 3 种类型 = 共 6 条",
        ],
        "notes": (
            "总结一下迷宫的安排："
            "「简单」难度下有 3 个迷宫，分别对应刚才说的盲走、旋律路线、带错误提醒这三种；"
            "「复杂」难度下也有同样的 3 个迷宫，只是迷宫本身更难一些。"
            "所以是 2 乘以 3，一共 6 条迷宫需要通关。"
            "我们会按顺序带你一个一个完成，不用着急。"
        ),
    },
    {
        "title": "走迷宫：你需要知道的",
        "bullets": [
            "设备：头戴式耳机 + 电脑（我们准备好）",
            "操作：键盘 ↑ ↓ ← → 四个方向键",
            "不用看屏幕，闭着眼睛听声音走就行",
        ],
        "notes": (
            "走迷宫时，我们会给你准备好头戴式耳机和电脑，你不需要自己带设备。"
            "操作非常简单：用键盘上的上下左右四个方向键控制移动。"
            "不需要看电脑屏幕，闭着眼睛、专心听声音就可以。"
            "如果某个方向走不通，换另一个方向试试；"
            "到达终点时会有明显的提示音告诉你成功了。"
        ),
        "link": "https://yukting-tse.github.io/HierarchyTone/Prototype/",
    },
    {
        "title": "第二部分：体验新的读屏软件",
        "bullets": [
            "读屏软件：帮你看手机/电脑上的内容",
            "现有读屏：点到哪里就念哪里",
            "但每次都要记住整个界面的结构",
        ],
        "notes": (
            "第二部分是体验一种新的读屏软件设计。"
            "如果你用过手机上的读屏功能，应该知道：手指滑到哪个按钮，它就会念出那个按钮的名字。"
            "这很方便，但问题是——每次打开一个新 App，你都要从头记住整个界面的布局，"
            "哪些是重要的、哪些是次要的，光靠听文字很难一下子分清楚。"
            "所以我们想试试：用不同的音高来告诉你信息的重要程度。"
        ),
    },
    {
        "title": "音高 = 信息的重要程度",
        "bullets": [
            "重要选项（如导航栏「首页」）：",
            "  「叮」（较高音）+「首页」",
            "普通内容（如图片说明文字）：",
            "  「咚」（较低音）+ 文字介绍",
        ],
        "notes": (
            "新的设计思路是这样的："
            "当你滑到比较重要的选项，比如导航栏里的「首页」，"
            "会先听到一声较高的「叮」，然后念出「首页」——高音代表「这是重要的地方」。"
            "当你滑到普通的图片或说明文字时，"
            "会先听到一声较低的「咚」，再念出文字内容——低音代表「这是普通信息」。"
            "这样你不用记住整个界面，只要听音高高低，就能大概知道现在摸到的是不是重要选项。"
        ),
    },
    {
        "title": "怎么操作这个演示？",
        "bullets": [
            "在电脑上测试，模拟手机滑动的用法",
            "鼠标移到哪个选项，就自动播报",
            "和现有读屏一样，不用学新功能",
        ],
        "notes": (
            "为了方便大家使用，这个演示结合了现有读屏软件的操作习惯。"
            "我们在电脑上做成了模拟手机操作的方式："
            "用鼠标移动到你想了解的选项上，系统就会自动播报那个选项的内容。"
            "你不需要提前学习什么新功能，用法和平时用手机读屏差不多，"
            "只是多了一种「叮」和「咚」的音高提示。"
        ),
        "link": "https://yukting-tse.github.io/HierarchyTone/",
    },
    {
        "title": "你会体验到什么？",
        "bullets": [
            "手机常用软件的功能（模拟界面）",
            "浏览网页的测试",
            "感受音高能不能帮你更快找到重要信息",
        ],
        "notes": (
            "在这个演示里，除了模拟手机常用软件的功能界面，"
            "还有浏览网页的测试场景。"
            "你可以自由滑动、探索不同的选项，"
            "感受一下：有了「叮」和「咚」的音高区别之后，"
            "找重要按钮、分清楚导航和内容，是不是比以前更容易一些。"
            "你的真实感受对我们非常重要。"
        ),
    },
    {
        "title": "实验结束后",
        "bullets": [
            "我们会口头聊聊：",
            "  · 用起来感觉怎么样？",
            "  · 有没有想改进的地方？",
            "  · 最喜欢 / 最不习惯的是哪部分？",
            "感谢参与！准备了小零食和饮料",
        ],
        "notes": (
            "两个实验都做完之后，我们会和你口头聊聊问卷里的几个问题，"
            "比如：整体用起来感觉怎么样？有没有什么地方觉得不好用、希望我们改进？"
            "你最喜欢哪个部分？哪个部分最不习惯？"
            "没有标准答案，怎么想就怎么说，你的意见会帮助我们做得更好。"
            "非常感谢你愿意花时间参加！我们还准备了一些小零食和饮料，"
            "实验结束后可以自取。再次谢谢你的参与！"
        ),
    },
    {
        "title": "准备好了吗？",
        "bullets": [
            "有问题现在可以提问",
            "准备好了我们就从走迷宫开始",
            "放轻松，像玩游戏一样就好",
        ],
        "notes": (
            "好，介绍就到这里。大家有什么问题现在可以提问。"
            "如果没有问题，我们就从第一部分的走迷宫开始。"
            "放轻松，就当玩游戏一样，重要的是你的真实感受。"
            "我们随时在旁边，有需要就叫我。"
        ),
    },
]


def set_slide_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title(slide, text):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.45), Inches(12.1), Inches(1.2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = ACCENT
    p.alignment = PP_ALIGN.LEFT
    return box


def add_bullets(slide, bullets, start_top=1.75):
    box = slide.shapes.add_textbox(Inches(0.75), Inches(start_top), Inches(11.8), Inches(5.0))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP

    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.font.size = Pt(28 if not bullet.startswith("  ") else 24)
        p.font.color.rgb = WHITE if not bullet.startswith("  ") else LIGHT_GRAY
        p.space_after = Pt(14)
        p.level = 0 if not bullet.startswith("  ") else 1
    return box


def add_link_footer(slide, url):
    box = slide.shapes.add_textbox(Inches(0.75), Inches(6.55), Inches(11.8), Inches(0.55))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = f"链接：{url}"
    p.font.size = Pt(16)
    p.font.color.rgb = SUBTLE


def add_speaker_notes(slide, notes):
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = notes


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    for data in SLIDES:
        slide = prs.slides.add_slide(blank_layout)
        set_slide_bg(slide, BG_DARK)
        add_title(slide, data["title"])
        add_bullets(slide, data["bullets"])
        if data.get("link"):
            add_link_footer(slide, data["link"])
        add_speaker_notes(slide, data["notes"])

    out = "/workspace/听觉实验说明_视障青少年.pptx"
    prs.save(out)
    print(f"Saved: {out}")
    return out


if __name__ == "__main__":
    build()
